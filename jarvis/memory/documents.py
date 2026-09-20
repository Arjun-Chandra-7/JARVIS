"""Find the file you half remember, by what was in it.

Jarvis could read a file it was handed the path to, and list a directory, and read a project. All
three want you to already know where the thing is. There was no way to answer "find that PDF
about the hackathon", which is how people actually look for their own documents — by content,
weeks later, having forgotten the filename entirely.

The vault is Jarvis's memory, not yours. This is yours.

Text first, layout never
------------------------
Everything here reduces a document to plain text and throws the rest away. That is the right
trade for search: nobody finds a file by its margins. It also keeps the extractors cheap and
swappable, which matters because every one of them is optional — `pdftotext` when poppler is
installed (it is, and it is the fastest), `pypdf` when it is not, `python-docx` for Word. Missing
all of them degrades to the plain-text formats rather than to an exception.

Never the whole disk
--------------------
Indexing a home directory means indexing browser caches, node_modules, virtualenvs and every
model checkpoint on the machine — gigabytes of noise around a few hundred real documents. So the
roots are named, the extensions are named, and anything enormous is skipped rather than read.

What ~/Downloads actually contains
----------------------------------
Worth saying plainly, because it was true on the machine this was written for: Downloads holds
identity documents. Scanned ID, bank letters, medical results — the things a browser drops there
and nobody files. Indexing that folder means their text is on disk in an index, and a question
that happens to match will read it back.

Nothing here tries to detect that automatically. Filename heuristics for "is this sensitive" are
a denylist, and a denylist that is wrong once is worse than no promise at all. Instead the roots
are configurable and named — `JARVIS_DOCUMENT_ROOTS` — so a person who does not want Downloads
indexed can say so, and a person who does has been told what that means.
"""

from __future__ import annotations

import os
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Optional

# Where a person's documents actually live. Overridable, because not everyone files things the
# same way, and Jarvis should not have an opinion about that.
DEFAULT_ROOTS = ("~/Documents", "~/Downloads", "~/Desktop")

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".org", ".csv", ".json", ".yaml", ".yml"}
PDF_SUFFIXES = {".pdf"}
WORD_SUFFIXES = {".docx"}

# Past this a plain-text file is a log, a dump or a mistake, not something anybody reads.
MAX_TEXT_BYTES = 12 * 1024 * 1024

# PDFs and Word files need their own, much larger ceiling, because their size says nothing about
# how much text is in them — it is almost entirely images. Found by testing rather than guessed:
# a perfectly ordinary illustrated guide in Downloads is 13 MB and extracts to a few pages of
# words, and a shared limit threw it out as if it were a dataset.
MAX_BINARY_BYTES = 120 * 1024 * 1024

# Directories that are never somebody's documents, however deep the walk goes.
SKIP_DIRS = {
    ".git", ".svn", "node_modules", "__pycache__", ".venv", "venv", ".cache", ".local",
    "site-packages", ".mypy_cache", ".pytest_cache", "dist", "build", ".next", "target",
    "snap", ".config", ".mozilla", ".steam", "Trash",
}


@dataclass
class Document:
    path: Path
    text: str
    kind: str

    @property
    def name(self) -> str:
        return self.path.name


def roots() -> list[Path]:
    configured = os.environ.get("JARVIS_DOCUMENT_ROOTS", "")
    names = [n for n in configured.split(":") if n.strip()] or list(DEFAULT_ROOTS)
    return [Path(os.path.expanduser(n)) for n in names]


def _pdf_text(path: Path) -> str:
    # poppler's pdftotext first: it is already on this machine, it is faster than anything in
    # Python, and it handles the malformed files pypdf gives up on.
    if shutil.which("pdftotext"):
        try:
            done = subprocess.run(["pdftotext", "-q", "-nopgbrk", str(path), "-"],
                                  capture_output=True, text=True, timeout=25)
            if done.stdout.strip():
                return done.stdout
        except Exception:  # noqa: BLE001
            pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:  # noqa: BLE001 — an unreadable PDF is skipped, not fatal
        return ""


def _word_text(path: Path) -> str:
    try:
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    except Exception:  # noqa: BLE001
        return ""


def extract(path: Path) -> Optional[Document]:
    """The readable text of one file, or None when there is none to be had."""
    suffix = path.suffix.lower()
    binary = suffix in PDF_SUFFIXES or suffix in WORD_SUFFIXES
    try:
        if not path.is_file():
            return None
        if path.stat().st_size > (MAX_BINARY_BYTES if binary else MAX_TEXT_BYTES):
            return None
    except OSError:
        return None

    if suffix in PDF_SUFFIXES:
        text, kind = _pdf_text(path), "pdf"
    elif suffix in WORD_SUFFIXES:
        text, kind = _word_text(path), "word"
    elif suffix in TEXT_SUFFIXES:
        try:
            text, kind = path.read_text(errors="ignore"), "text"
        except OSError:
            return None
    else:
        return None

    text = text.strip()
    return Document(path=path, text=text, kind=kind) if text else None


def walk(where: Optional[list[Path]] = None, limit: int = 4000) -> Iterator[Path]:
    """Candidate files under the configured roots, skipping the places that are never documents."""
    seen = 0
    for root in (where if where is not None else roots()):
        if not root.exists():
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            # Pruned in place so os.walk never descends — checking after the fact would still
            # cost the walk of a node_modules tree.
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
            for name in filenames:
                suffix = Path(name).suffix.lower()
                if suffix in TEXT_SUFFIXES or suffix in PDF_SUFFIXES or suffix in WORD_SUFFIXES:
                    yield Path(dirpath) / name
                    seen += 1
                    if seen >= limit:
                        return
